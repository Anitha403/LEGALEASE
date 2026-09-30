from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from utils.text import sanitize_text


ROOT = Path(__file__).resolve().parents[1]

LOGO_PATH = (
    ROOT /
    "assets" /
    "logo.png"
)


def format_txt(text: str) -> bytes:

    return sanitize_text(
        text
    ).encode("utf-8")


def _add_docx_logo(
    doc: Document
):

    if LOGO_PATH.exists():

        paragraph = doc.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run()

        run.add_picture(
            str(LOGO_PATH),
            width=Inches(1.2)
        )


def format_docx(
    text: str,
    doc_type: str = "Legal Document"
) -> bytes:

    doc = Document()


    section = doc.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)


    _add_docx_logo(doc)


    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    run = title.add_run(
        doc_type.upper()
    )

    run.bold = True

    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(16)


    for block in sanitize_text(
        text
    ).split("\n\n"):

        block = block.strip()

        if not block:
            continue


        lines = block.split("\n")

        first = lines[0].strip()


        if (
            first.isupper()
            and len(first) <= 100
            and not first.startswith("-")
        ):

            paragraph = doc.add_paragraph()

            run = paragraph.add_run(
                first
            )

            run.bold = True

            run.font.name = (
                "Times New Roman"
            )

            run.font.size = Pt(12)

            continue


        for line in lines:

            paragraph = doc.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(6)

            run = paragraph.add_run(
                line
            )

            run.font.name = (
                "Times New Roman"
            )

            run.font.size = Pt(11)


    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    footer_run = footer.add_run(
        "LegalEase - AI-generated draft. "
        "Review with a qualified legal professional."
    )

    footer_run.font.name = (
        "Times New Roman"
    )

    footer_run.font.size = Pt(8)


    output = BytesIO()

    doc.save(output)

    return output.getvalue()


class LegalPDF(FPDF):

    def __init__(
        self,
        doc_type: str
    ):

        super().__init__()

        self.doc_type = doc_type


    def header(self):

        if LOGO_PATH.exists():

            self.image(
                str(LOGO_PATH),
                x=92,
                y=8,
                w=26
            )

            self.set_y(37)

        else:

            self.set_y(12)


        self.set_font(
            "Helvetica",
            "B",
            12
        )


        self.cell(
            0,
            8,
            self.doc_type.upper(),
            align="C"
        )


        self.ln(10)


    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            size=7
        )


        self.cell(
            0,
            8,
            "LegalEase - AI-generated draft for review",
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str = "Legal Document"
) -> bytes:

    pdf = LegalPDF(
        doc_type
    )


    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )


    pdf.add_page()


    pdf.set_font(
        "Helvetica",
        size=10
    )


    safe_text = sanitize_text(
        text
    )


    for block in safe_text.split(
        "\n\n"
    ):

        block = block.strip()

        if not block:
            continue


        lines = block.split("\n")

        first = lines[0].strip()


        if (
            first.isupper()
            and len(first) <= 100
            and not first.startswith("-")
        ):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )


            pdf.multi_cell(
                0,
                7,
                first
            )


            pdf.ln(2)


            pdf.set_font(
                "Helvetica",
                size=10
            )

            continue


        for line in lines:

            if line.startswith("- "):

                pdf.multi_cell(
                    0,
                    6,
                    "- " + line[2:]
                )

            else:

                pdf.multi_cell(
                    0,
                    6,
                    line
                )


        pdf.ln(2)


    return bytes(
        pdf.output()
    )